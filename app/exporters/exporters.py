from __future__ import annotations
from pathlib import Path
from docx import Document
from app.core.models import TranscriptResult, Segment


def _time_srt(seconds: float) -> str:
    ms = max(0, int(round(seconds * 1000)))
    h, rem = divmod(ms, 3600000)
    m, rem = divmod(rem, 60000)
    s, ms = divmod(rem, 1000)
    return f'{h:02}:{m:02}:{s:02},{ms:03}'


def _time_vtt(seconds: float) -> str:
    return _time_srt(seconds).replace(',', '.')


def _line(seg: Segment) -> str:
    who = f'{seg.speaker}：' if seg.speaker else ''
    return f'{who}{seg.text}'


def export_all(result: TranscriptResult, base_path: str, formats: list[str], note: str | None = None) -> list[str]:
    """Export a transcript to user-selected formats.

    `note` is used for partial/aborted jobs. It is included in TXT, DOCX and VTT
    without corrupting SRT timing structure. The filename itself also carries the
    中止版 marker, so SRT remains standards-friendly.
    """
    base = Path(base_path)
    base.parent.mkdir(parents=True, exist_ok=True)
    out = []
    if 'txt' in formats:
        p = base.with_suffix('.txt')
        body = result.text or '\n'.join(_line(x) for x in result.segments)
        if note:
            body = note + '\n\n' + body
        p.write_text(body, encoding='utf-8-sig')
        out.append(str(p))
    if 'srt' in formats and result.segments:
        p = base.with_suffix('.srt')
        blocks = []
        for i, s in enumerate(result.segments, 1):
            end = s.end if s.end > s.start else s.start + 2.0
            blocks.append(f'{i}\n{_time_srt(s.start)} --> {_time_srt(end)}\n{_line(s)}')
        p.write_text('\n\n'.join(blocks), encoding='utf-8-sig')
        out.append(str(p))
    if 'vtt' in formats and result.segments:
        p = base.with_suffix('.vtt')
        blocks = ['WEBVTT\n']
        if note:
            blocks.append(f'NOTE {note}\n')
        for s in result.segments:
            end = s.end if s.end > s.start else s.start + 2.0
            blocks.append(f'{_time_vtt(s.start)} --> {_time_vtt(end)}\n{_line(s)}\n')
        p.write_text('\n'.join(blocks), encoding='utf-8-sig')
        out.append(str(p))
    if 'docx' in formats:
        p = base.with_suffix('.docx')
        doc = Document()
        doc.add_heading(base.stem, level=1)
        doc.add_paragraph(f'辨識引擎：{result.engine}')
        if result.language:
            doc.add_paragraph(f'語言：{result.language}')
        if note:
            doc.add_paragraph(note)
        doc.add_heading('逐字稿', level=2)
        if result.segments:
            for s in result.segments:
                t = f'[{_time_vtt(s.start)[:-4]}] '
                if s.speaker:
                    t += f'{s.speaker}：'
                doc.add_paragraph(t + s.text)
        else:
            doc.add_paragraph(result.text)
        doc.save(p)
        out.append(str(p))
    return out
